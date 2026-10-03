"""Document processing: parsing and semantic chunking."""

import os
import logging
from typing import Optional
import fitz  # PyMuPDF
from docx import Document as DocxDocument
from llama_index.core.schema import TextNode
from llama_index.core.node_parser import SemanticSplitterNodeParser
from llama_index.core import Document
from app.embedding_service import get_embed_model

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def parse_pdf(file_path: str) -> list[dict]:
    """Extract text from PDF using PyMuPDF, page by page."""
    pages = []
    doc = fitz.open(file_path)
    for page_num, page in enumerate(doc):
        text = page.get_text("text").strip()
        if text:
            pages.append({"text": text, "page": page_num + 1})
    doc.close()
    return pages


def _format_docx_table(table) -> str:
    """Convert a python-docx Table object to Markdown table syntax."""
    rows = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        if any(cells):
            rows.append("| " + " | ".join(cells) + " |")
    if not rows:
        return ""
    if len(rows) > 1:
        header = rows[0]
        divider = "| " + " | ".join(["---"] * len(table.rows[0].cells)) + " |"
        return "\n".join([header, divider] + rows[1:])
    return "\n".join(rows)


def parse_docx(file_path: str) -> list[dict]:
    """Extract text and tables from DOCX file."""
    doc = DocxDocument(file_path)
    content_blocks = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            content_blocks.append(text)
    for table in doc.tables:
        table_md = _format_docx_table(table)
        if table_md:
            content_blocks.append(table_md)
    full_text = "\n\n".join(content_blocks)
    if full_text:
        return [{"text": full_text, "page": 1}]
    return []


def parse_text(file_path: str) -> list[dict]:
    """Extract text from TXT/MD files."""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read().strip()
    if text:
        return [{"text": text, "page": 1}]
    return []


def parse_document(file_path: str) -> list[dict]:
    """Parse a document based on its extension."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext == ".docx":
        return parse_docx(file_path)
    elif ext in (".txt", ".md"):
        return parse_text(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def _clamp_node_length(
    nodes: list[TextNode], max_chars: int = 1800, overlap: int = 150
) -> list[TextNode]:
    """
    Ensure no node exceeds max_chars (~450 tokens) to prevent silent truncation
    in embedding (512 tokens) and cross-encoder reranking models.
    """
    clamped_nodes = []
    for node in nodes:
        text = node.get_content()
        if len(text) <= max_chars:
            clamped_nodes.append(node)
            continue

        start = 0
        sub_idx = 0
        while start < len(text):
            end = min(start + max_chars, len(text))
            if end < len(text):
                # Try breaking at newline or sentence boundary
                break_point = max(text.rfind("\n", start, end), text.rfind(". ", start, end))
                if break_point > start + max_chars // 2:
                    end = break_point + 1
            chunk_text = text[start:end].strip()
            if chunk_text:
                new_meta = node.metadata.copy()
                new_meta["sub_chunk"] = sub_idx
                sub_node = TextNode(text=chunk_text, metadata=new_meta)
                clamped_nodes.append(sub_node)
                sub_idx += 1
            start = end - overlap if end < len(text) else end
    return clamped_nodes


def semantic_chunk_document(
    file_path: str,
    doc_name: str,
    buffer_size: int = 1,
    breakpoint_percentile_threshold: int = 95,
) -> list[TextNode]:
    """
    Parse a document and split it into semantically coherent chunks.

    Uses LlamaIndex SemanticSplitterNodeParser which groups sentences
    by embedding similarity — creating chunks that preserve meaning,
    followed by length clamping to prevent context window overflow.
    """
    logger.info(f"Processing document: {doc_name}")

    # Parse the document into pages
    pages = parse_document(file_path)
    if not pages:
        raise ValueError(f"No text could be extracted from {doc_name}")

    # Convert to LlamaIndex Documents
    llama_docs = []
    for page_info in pages:
        doc = Document(
            text=page_info["text"],
            metadata={"doc_name": doc_name, "page": page_info["page"]},
        )
        llama_docs.append(doc)

    # Semantic chunking
    embed_model = get_embed_model()
    splitter = SemanticSplitterNodeParser(
        buffer_size=buffer_size,
        breakpoint_percentile_threshold=breakpoint_percentile_threshold,
        embed_model=embed_model,
    )

    nodes = splitter.get_nodes_from_documents(llama_docs)
    nodes = _clamp_node_length(nodes, max_chars=1800, overlap=150)

    # Add chunk index metadata
    for idx, node in enumerate(nodes):
        node.metadata["chunk_index"] = idx
        node.metadata["doc_name"] = doc_name

    logger.info(f"Created {len(nodes)} semantic chunks from {doc_name}")
    return nodes
